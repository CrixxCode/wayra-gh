"""
Exportacion completa de reportes a PDF y Excel (auditoria, Bloque 11 #6-7 y Bloque 9 #11).

Antes "Exportar PDF" abria un popup con cuatro lineas de KPIs (y fallaba en silencio si el
navegador lo bloqueaba), y no habia Excel. Aqui se recorre el payload del reporte —el mismo que
pinta la pantalla— y se arma el documento: periodo, indicadores y, por cada seccion, su grafico
y su tabla. Es generico a proposito: todos los reportes comparten la forma `filters` +
`kpis`/`summary` + listas de filas, asi que un reporte nuevo se exporta sin codigo extra.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

# --------------------------------------------------------------------- etiquetas

LABELS = {
    # secciones
    "income_vs_profit_chart": "Ingresos y utilidad por mes",
    "payment_methods": "Metodos de pago",
    "weekly_occupancy": "Ocupacion semanal",
    "top_guests": "Huespedes principales",
    "monthly_income_vs_expenses": "Ingresos y egresos por mes",
    "monthly_net_profit": "Utilidad neta por mes",
    "payment_breakdown": "Desglose por metodo de pago",
    "guest_origin": "Origen de los huespedes",
    "monthly_occupancy_rate": "Ocupacion por mes",
    "occupied_rooms_by_month": "Habitaciones ocupadas por mes",
    "room_type_performance": "Rendimiento por tipo de habitacion",
    "income_by_category": "Ingresos por categoria",
    "transactions_by_category": "Transacciones por categoria",
    "category_detail": "Detalle por categoria",
    "daily_rows": "Ingresos por dia",
    "method_rows": "Ingresos por metodo",
    "expenses": "Egresos",
    "by_category": "Egresos por categoria",
    # indicadores
    "annual_income": "Ingresos del periodo",
    "net_profit": "Utilidad neta",
    "average_occupancy": "Ocupacion promedio (%)",
    "revpar": "RevPAR",
    "gross_income": "Ingresos brutos",
    "total_expenses": "Egresos totales",
    "net_margin": "Margen neto (%)",
    "occupancy_peak": "Pico de ocupacion (%)",
    "average_stay": "Estadia promedio (noches)",
    "total_guests": "Huespedes",
    "service_income": "Ingresos por servicios",
    "transactions": "Transacciones",
    "average_ticket": "Ticket promedio",
    "top_category": "Categoria principal",
    "total_transactions": "Transacciones",
    "active_transactions": "Transacciones activas",
    "total_collected": "Total recaudado",
    "today_collected": "Recaudado hoy",
    "month_collected": "Recaudado en el mes",
    "total_amount": "Total",
    "count": "Cantidad",
    # columnas
    "month": "Mes",
    "week": "Semana",
    "income": "Ingresos",
    "profit": "Utilidad",
    "expenses_amount": "Egresos",
    "value": "Valor",
    "method": "Metodo",
    "method_label": "Metodo",
    "amount": "Monto",
    "pct": "%",
    "amount_pct": "% del monto",
    "transactions_pct": "% de transacciones",
    "share_pct": "Participacion (%)",
    "share_percent": "Participacion (%)",
    "trend_pct": "Tendencia (%)",
    "occupied_rooms": "Habitaciones ocupadas",
    "occupancy_rate_pct": "Ocupacion (%)",
    "occupancy_pct": "Ocupacion (%)",
    "rooms": "Habitaciones",
    "room_type": "Tipo de habitacion",
    "avg_stay": "Estadia promedio",
    "guest_name": "Huesped",
    "country": "Pais",
    "stays": "Estadias",
    "nights": "Noches",
    "total_spent": "Gasto total",
    "segment": "Segmento",
    "category": "Categoria",
    "date_label": "Fecha",
    "inactive_transactions": "Transacciones inactivas",
    "top_method": "Metodo principal",
    "top_guest": "Huesped principal",
    "expense_date": "Fecha",
    "concept": "Concepto",
    "supplier_name": "Proveedor",
    "reference": "Referencia",
    "payment_method": "Metodo de pago",
}

# Columnas internas que no aportan al lector.
HIDDEN_COLUMNS = {"date_key", "method_key", "id"}
# Secciones duplicadas en el payload (la pantalla las usa para dos graficos).
HIDDEN_SECTIONS = {"by_room_type"}
# Columnas que hacen de etiqueta del eje en los graficos.
LABEL_COLUMNS = ("month", "week", "method", "method_label", "category", "room_type", "country", "date_label", "guest_name", "name")


def label(key: str) -> str:
    return LABELS.get(key) or key.replace("_", " ").capitalize()


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _format_value(value) -> str:
    if value is None:
        return "-"
    if _is_number(value):
        if isinstance(value, float) and not value.is_integer():
            return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{int(value):,}".replace(",", ".")
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _indicator_rows(payload: dict) -> list[tuple[str, Any, str]]:
    """(indicador, valor, variacion) a partir de `kpis` o `summary`."""
    rows: list[tuple[str, Any, str]] = []
    for block_key in ("kpis", "summary"):
        block = payload.get(block_key) or {}
        for key, item in block.items():
            if isinstance(item, dict):
                value = item.get("value", item.get("amount"))
                extra = ""
                for variation_key, suffix in (
                    ("variation_pct", " %"),
                    ("variation_points", " pts"),
                    ("variation_value", ""),
                    ("variation_nights", " noches"),
                ):
                    if item.get(variation_key) is not None:
                        extra = f"{_format_value(item[variation_key])}{suffix}"
                        break
                if item.get("month"):
                    extra = str(item["month"])
                if item.get("name"):
                    extra = str(item["name"])
                rows.append((label(key), value, extra))
            else:
                rows.append((label(key), item, ""))
    return rows


def _sections(payload: dict) -> list[tuple[str, list[str], list[dict]]]:
    """(titulo, columnas, filas) de cada lista de filas del payload."""
    sections = []
    for key, value in payload.items():
        if key in HIDDEN_SECTIONS or not isinstance(value, list) or not value:
            continue
        if not all(isinstance(row, dict) for row in value):
            continue
        columns = [column for column in value[0].keys() if column not in HIDDEN_COLUMNS]
        sections.append((label(key), columns, value))
    return sections


def _chart_spec(columns: list[str], rows: list[dict]):
    """Columna de etiquetas y series numericas, si la seccion se puede graficar."""
    label_column = next((column for column in LABEL_COLUMNS if column in columns), None)
    numeric = [column for column in columns if rows and _is_number(rows[0].get(column))]
    if not label_column or not numeric or len(rows) < 2:
        return None
    # Montos y porcentajes juntos no comparten eje: se grafica el primer grupo homogeneo.
    pct_columns = [column for column in numeric if "pct" in column or "percent" in column]
    series = pct_columns if len(pct_columns) == len(numeric) else [c for c in numeric if c not in pct_columns]
    return label_column, series[:3]


def _period_text(payload: dict) -> str:
    filters = payload.get("filters") or {}
    start, end = filters.get("start_date"), filters.get("end_date")
    if start and end:
        return f"Periodo: {_format_value(start)} a {_format_value(end)}"
    return ""


# ------------------------------------------------------------------------- Excel


def render_xlsx(title: str, payload: dict, *, hotel_name: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F1F41")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append([title])
    summary["A1"].font = Font(bold=True, size=14)
    if hotel_name:
        summary.append([hotel_name])
    period = _period_text(payload)
    if period:
        summary.append([period])
    summary.append([])
    indicators = _indicator_rows(payload)
    if indicators:
        summary.append(["Indicador", "Valor", "Variacion / detalle"])
        for cell in summary[summary.max_row]:
            cell.font, cell.fill = header_font, header_fill
        for name, value, extra in indicators:
            summary.append([name, value if _is_number(value) else _format_value(value), extra])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 24

    used_names = {"Resumen"}
    for section_title, columns, rows in _sections(payload):
        name = section_title[:28]
        suffix = 2
        while name in used_names:
            name = f"{section_title[:25]} {suffix}"
            suffix += 1
        used_names.add(name)
        sheet = workbook.create_sheet(name)
        sheet.append([label(column) for column in columns])
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append([
                row.get(column) if _is_number(row.get(column)) else _format_value(row.get(column))
                for column in columns
            ])
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(14, len(label(column)) + 4)
            if rows and _is_number(rows[0].get(column)):
                for cell in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                    cell[0].number_format = "0.00" if ("pct" in column or "percent" in column) else "#,##0.00"

        spec = _chart_spec(columns, rows)
        if spec:
            label_column, series = spec
            chart = BarChart()
            chart.title = section_title
            chart.height, chart.width = 8, 18
            for column in series:
                col_index = columns.index(column) + 1
                data = Reference(sheet, min_col=col_index, min_row=1, max_row=len(rows) + 1)
                chart.add_data(data, titles_from_data=True)
            categories = Reference(
                sheet, min_col=columns.index(label_column) + 1, min_row=2, max_row=len(rows) + 1
            )
            chart.set_categories(categories)
            sheet.add_chart(chart, f"{get_column_letter(len(columns) + 2)}2")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF


def render_pdf(title: str, payload: dict, *, hotel_name: str = "") -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    brand = colors.HexColor("#0f1f41")
    palette = [colors.HexColor("#2a78d6"), colors.HexColor("#1baf7a"), colors.HexColor("#e34948")]
    styles = getSampleStyleSheet()
    buffer = BytesIO()
    page_width, _ = landscape(A4)
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=title,
    )
    usable_width = page_width - 3 * cm
    story = [Paragraph(title, styles["Title"])]
    subtitle = " · ".join(part for part in (hotel_name, _period_text(payload)) if part)
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 0.4 * cm))

    table_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), brand),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d7deea")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fbff")]),
        ]
    )

    def styled_table(data, col_widths, numeric_columns):
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(table_style)
        # Solo las cifras van a la derecha; el texto (pais, segmento...) se lee mejor a la izquierda.
        table.setStyle(TableStyle([("ALIGN", (index, 1), (index, -1), "RIGHT") for index in numeric_columns]))
        return table

    indicators = _indicator_rows(payload)
    if indicators:
        story.append(Paragraph("Indicadores", styles["Heading2"]))
        data = [["Indicador", "Valor", "Variacion / detalle"]] + [
            [name, _format_value(value), extra] for name, value, extra in indicators
        ]
        table = styled_table(
            data, [usable_width * 0.45, usable_width * 0.25, usable_width * 0.30], numeric_columns=[1]
        )
        story += [table, Spacer(1, 0.5 * cm)]

    for section_title, columns, rows in _sections(payload):
        # El titulo viaja con su grafico (o con su tabla corta): no queda huerfano al pie.
        block = [Paragraph(section_title, styles["Heading2"])]
        spec = _chart_spec(columns, rows)
        if spec:
            label_column, series = spec
            drawing = Drawing(usable_width, 6 * cm)
            chart = VerticalBarChart()
            chart.x, chart.y = 1.5 * cm, 0.8 * cm
            chart.width, chart.height = usable_width - 2.5 * cm, 4.6 * cm
            chart.data = [[float(row.get(column) or 0) for row in rows] for column in series]
            chart.categoryAxis.categoryNames = [str(row.get(label_column) or "")[:12] for row in rows]
            chart.categoryAxis.labels.fontSize = 6
            chart.categoryAxis.labels.angle = 30 if len(rows) > 8 else 0
            chart.valueAxis.labels.fontSize = 6
            chart.valueAxis.labelTextFormat = lambda value: _format_value(float(value))
            chart.valueAxis.valueMin = min(0, min((min(serie) for serie in chart.data), default=0))
            for index in range(len(series)):
                chart.bars[index].fillColor = palette[index % len(palette)]
            drawing.add(chart)
            block.append(drawing)
            legend = " · ".join(label(column) for column in series)
            block.append(Paragraph(f"<font size=7>{legend}</font>", styles["Normal"]))
        data = [[label(column) for column in columns]] + [
            [_format_value(row.get(column)) for column in columns] for row in rows
        ]
        numeric_columns = [index for index, column in enumerate(columns) if _is_number(rows[0].get(column))]
        table = styled_table(data, [usable_width / len(columns)] * len(columns), numeric_columns)
        if len(rows) <= 12:
            story += [KeepTogether(block + [table]), Spacer(1, 0.5 * cm)]
        else:
            story += [KeepTogether(block), table, Spacer(1, 0.5 * cm)]

    if not indicators and not _sections(payload):
        story.append(Paragraph("No hay datos para el periodo seleccionado.", styles["Normal"]))

    document.build(story)
    return buffer.getvalue()


def render(fmt: str, title: str, payload: dict, *, hotel_name: str = "") -> tuple[bytes, str, str]:
    """(contenido, content-type, extension)."""
    if fmt == "xlsx":
        return (
            render_xlsx(title, payload, hotel_name=hotel_name),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    return render_pdf(title, payload, hotel_name=hotel_name), "application/pdf", "pdf"
